#!/usr/bin/env python3
"""
Build hrt-app/data/wallet-catalog.js from the My Wallet card catalog sheet.

Source of truth: the Google Sheet "HRT-My-Wallet-card-catalog".
Three ways to feed it the sheet:

    python3 tools/build_wallet_catalog.py ~/Downloads/HRT-My-Wallet-card-catalog.xlsx
    python3 tools/build_wallet_catalog.py --csv-dir folder/   (Cards.csv, Benefits.csv, ...)
    python3 tools/build_wallet_catalog.py --urls Cards=<url> Benefits=<url> Earning=<url> Quarterly=<url>
        (the "Publish to web" CSV links; this is what the daily update uses)

Safety: every tab named in --require must be present, and if Cards, Benefits
or Earning would shrink below 80% of what the site has now, nothing is
written (a half-loaded sheet must never wipe the site). To remove many rows
on purpose, add --allow-shrink.

It refuses to write anything if the sheet has a problem, and says which row.

Rules it enforces (from the sheet's Read me tab):
  - card_id is unique and matches on both tabs
  - cadence and resets use the agreed words only
  - benefits marked "unverified" never reach members, so they are dropped here
No secrets live in the sheet or this script. Safe for the public repo.
"""
import csv
import io
import json
import re
import sys
import urllib.request
from datetime import date
from pathlib import Path


OUT = Path(__file__).resolve().parent.parent / "hrt-app" / "data" / "wallet-catalog.js"

CADENCES = {"monthly", "quarterly", "semiannual", "annual", "every4years", "per event", "per purchase"}
RESETS = {"calendar", "anniversary", "unclear", "application date", "n/a"}
CONFIDENCE = {"confirmed", "assumed", "unverified"}
UNITS = {"usd", "points", "passes", "certificates"}
FEE_BASIS = {"anniversary", "n/a"}
GROUPS = {"online", "everyday", "dining", "grocery", "gas", "ev", "transit", "streaming", "drugstores", "flights",
          "hotels", "car-rental", "travel", "foreign", "entertainment", "business", "promo", "rotating", "status"}
QUARTER_RE = re.compile(r"^\d{4}-Q[1-4]$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
ID_RE = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")


class SheetError(Exception):
    pass


def rows(table):
    """Yield (row_number, dict) for each non-empty row under the header row.
    table is a list of row tuples, header first."""
    it = iter(table)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    for n, raw in enumerate(it, start=2):
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        yield n, {h: ("" if v is None else v) for h, v in zip(header, raw) if h}


def text(v):
    if isinstance(v, (date,)):
        return v.isoformat()[:10]
    return str(v).strip()


def money(v, where):
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = text(v).replace("$", "").replace(",", "")
    try:
        return round(float(s), 2)
    except ValueError:
        raise SheetError(f"{where}: '{v}' is not a dollar amount")


def number(v, where):
    if isinstance(v, (int, float)):
        return float(v)
    s = text(v).replace(",", "").replace("$", "")
    try:
        return float(s)
    except ValueError:
        raise SheetError(f"{where}: '{v}' is not a number")


def checked_date(v, where):
    s = text(v)
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        raise SheetError(f"{where}: checked date '{v}' should look like 2026-10-13")
    return s


# Card art: an https link only. No spaces, quotes, brackets or backslashes, so the
# link is always safe inside an <img src="..."> attribute.
ART_RE = re.compile(r'^https://[^\s"\'<>()\\]+$')


def build_cards(ws):
    cards, seen = [], set()
    for n, r in rows(ws):
        where = f"Cards row {n}"
        cid = text(r.get("card_id", ""))
        if not ID_RE.match(cid):
            raise SheetError(f"{where}: card_id '{cid}' must be lowercase-with-dashes")
        if cid in seen:
            raise SheetError(f"{where}: card_id '{cid}' appears twice")
        seen.add(cid)
        basis = text(r.get("fee_basis", "")).lower()
        if basis not in FEE_BASIS:
            raise SheetError(f"{where}: fee_basis '{basis}' must be one of {sorted(FEE_BASIS)}")
        name = text(r.get("card_name", ""))
        if not name:
            raise SheetError(f"{where}: card_name is empty")
        art = text(r.get("card_art", ""))
        if art and not ART_RE.match(art):
            raise SheetError(f"{where}: card_art must be an https image link with no spaces or quotes")
        card = {
            "id": cid,
            "issuer": text(r.get("issuer", "")),
            "name": name,
            "network": text(r.get("network", "")),
            "fee": money(r.get("annual_fee", ""), where),
            "feeBasis": basis,
            "note": text(r.get("notes", "")),
            "source": text(r.get("source_url", "")),
            "checked": checked_date(r.get("checked", ""), where),
        }
        if art:
            card["art"] = art
        ff = text(r.get("foreign_fee", "")).replace("%", "").strip()
        if ff:
            try:
                ffv = float(ff)
            except ValueError:
                raise SheetError(f"{where}: foreign_fee '{ff}' must be a percent like 3 or 0")
            if not 0 <= ffv <= 10:
                raise SheetError(f"{where}: foreign_fee must be between 0 and 10")
            if ffv:
                card["ftf"] = ffv
        cards.append(card)
    if not cards:
        raise SheetError("Cards tab has no cards")
    return cards


def build_benefits(ws, card_ids):
    out, dropped, seen = [], 0, set()
    for n, r in rows(ws):
        where = f"Benefits row {n}"
        cid = text(r.get("card_id", ""))
        if cid not in card_ids:
            raise SheetError(f"{where}: card_id '{cid}' is not on the Cards tab")
        bid = text(r.get("benefit_id", ""))
        if not ID_RE.match(bid):
            raise SheetError(f"{where}: benefit_id '{bid}' must be lowercase-with-dashes")
        if (cid, bid) in seen:
            raise SheetError(f"{where}: {cid}/{bid} appears twice")
        seen.add((cid, bid))
        conf = text(r.get("confidence", "")).lower()
        if conf not in CONFIDENCE:
            raise SheetError(f"{where}: confidence '{conf}' must be one of {sorted(CONFIDENCE)}")
        if conf == "unverified":
            dropped += 1
            continue
        cadence = text(r.get("cadence", "")).lower()
        if cadence not in CADENCES:
            raise SheetError(f"{where}: cadence '{cadence}' must be one of {sorted(CADENCES)}")
        resets = text(r.get("resets", "")).lower()
        if resets not in RESETS:
            raise SheetError(f"{where}: resets '{resets}' must be one of {sorted(RESETS)}")
        unit = text(r.get("unit", "")).lower()
        if unit not in UNITS:
            raise SheetError(f"{where}: unit '{unit}' must be one of {sorted(UNITS)}")
        amt = r.get("amount_per_period", "")
        out.append({
            "card": cid,
            "id": bid,
            "label": text(r.get("label", "")),
            "type": text(r.get("type", "")),
            "unit": unit,
            "amount": money(amt, where) if unit == "usd" else number(amt, where),
            "perYear": number(r.get("periods_per_year", ""), where),
            "cadence": cadence,
            "resets": resets,
            "enroll": text(r.get("enrollment", "")).lower(),
            "spendGate": text(r.get("spend_gate", "")),
            "confidence": conf,
            "note": text(r.get("note", "")),
            "source": text(r.get("source_url", "")),
            "checked": checked_date(r.get("checked", ""), where),
        })
    return out, dropped


def blankish(v):
    """Research fills empty cells with words like 'None stated' or 'Any'. Treat them as empty."""
    t = text(v)
    return "" if t.lower() in {"", "none", "none stated", "none (unlimited)", "any", "card spend", "n/a"} else t


def build_earning(table, card_ids):
    out, dropped, seen = [], 0, set()
    for n, r in rows(table):
        where = f"Earning row {n}"
        cid = text(r.get("card_id", ""))
        if cid not in card_ids:
            raise SheetError(f"{where}: card_id '{cid}' is not on the Cards tab")
        cat = text(r.get("category", ""))
        if not cat:
            raise SheetError(f"{where}: category is empty")
        if (cid, cat.lower()) in seen:
            raise SheetError(f"{where}: {cid} has '{cat}' twice")
        seen.add((cid, cat.lower()))
        groups = text(r.get("group", "")).lower().split()
        bad = [g for g in groups if g not in GROUPS]
        if not groups or bad:
            raise SheetError(f"{where}: group '{' '.join(bad) or '(empty)'}' must be from {sorted(GROUPS)}")
        conf = text(r.get("confidence", "")).lower()
        if conf not in CONFIDENCE:
            raise SheetError(f"{where}: confidence '{conf}' must be one of {sorted(CONFIDENCE)}")
        if conf == "unverified":
            dropped += 1
            continue
        rate = number(r.get("rate", ""), where)
        if not (0 < rate <= 25):
            raise SheetError(f"{where}: rate {rate} looks wrong")
        out.append({
            "card": cid, "category": cat, "groups": groups, "rate": rate,
            "currency": text(r.get("currency", "")),
            "cap": blankish(r.get("cap", "")), "via": blankish(r.get("via", "")),
            "note": text(r.get("notes", "")), "confidence": conf,
            "source": text(r.get("source_url", "")),
            "checked": checked_date(r.get("checked", ""), where),
        })
    return out, dropped


def build_quarterly(table, card_ids):
    out, seen = [], set()
    for n, r in rows(table):
        where = f"Quarterly row {n}"
        cid = text(r.get("card_id", ""))
        if cid not in card_ids:
            raise SheetError(f"{where}: card_id '{cid}' is not on the Cards tab")
        q = text(r.get("quarter", ""))
        if not QUARTER_RE.match(q):
            raise SheetError(f"{where}: quarter '{q}' should look like 2026-Q4")
        if (cid, q) in seen:
            raise SheetError(f"{where}: {cid} {q} appears twice")
        seen.add((cid, q))
        dl = text(r.get("activation_deadline", ""))
        if dl and not DATE_RE.match(dl):
            raise SheetError(f"{where}: activation_deadline '{dl}' should look like 2026-12-14 or be empty")
        conf = text(r.get("confidence", "")).lower()
        if conf not in CONFIDENCE:
            raise SheetError(f"{where}: confidence '{conf}' must be one of {sorted(CONFIDENCE)}")
        if conf == "unverified":
            continue
        out.append({
            "card": cid, "quarter": q, "categories": text(r.get("categories", "")),
            "rate": number(r.get("rate", ""), where), "cap": number(r.get("spend_cap", "0") or 0, where),
            "deadline": dl, "note": text(r.get("notes", "")), "confidence": conf,
            "source": text(r.get("source_url", "")), "checked": checked_date(r.get("checked", ""), where),
        })
    return out


def no_em_dashes(obj):
    """House style: no em or en dashes reach members."""
    s = json.dumps(obj, ensure_ascii=False)
    if "—" in s or "–" in s:
        raise SheetError("A dash (em or en) is in the sheet text. House style says none in member-facing text.")


def tables_from_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("Needs openpyxl:  pip3 install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    return {name: list(wb[name].iter_rows(values_only=True)) for name in wb.sheetnames}


def tables_from_csv_text(texts):
    return {name: list(csv.reader(io.StringIO(t))) for name, t in texts.items()}


def tables_from_csv_dir(folder):
    texts = {}
    for f in Path(folder).glob("*.csv"):
        texts[f.stem] = f.read_text(encoding="utf-8-sig")
    return tables_from_csv_text(texts)


def tables_from_urls(pairs):
    texts = {}
    for pair in pairs:
        name, _, url = pair.partition("=")
        if not url.startswith("https://"):
            raise SheetError(f"'{pair}' should look like Cards=https://...")
        with urllib.request.urlopen(url, timeout=60) as r:
            body = r.read().decode("utf-8-sig")
        if body.lstrip().lower().startswith("<!doctype") or "<html" in body[:500].lower():
            raise SheetError(f"{name}: got a web page, not CSV. Is the tab published to the web as CSV?")
        texts[name] = body
    return tables_from_csv_text(texts)


EXPECTED_HEADERS = {
    "Cards": {"card_id", "issuer", "card_name", "annual_fee", "fee_basis", "checked"},
    "Benefits": {"card_id", "benefit_id", "label", "cadence", "resets", "confidence", "checked"},
    "Earning": {"card_id", "category", "group", "rate", "currency", "confidence", "checked"},
    "Quarterly": {"card_id", "quarter", "categories", "activation_deadline", "confidence", "checked"},
}


def check_headers(tables):
    """Google serves the FIRST published tab when an unpublished one is asked for,
    so a tab must prove it is the tab we think it is by its column headers."""
    for name, need in EXPECTED_HEADERS.items():
        if name not in tables or not tables[name]:
            continue
        got = {str(h).strip() for h in tables[name][0] if h is not None}
        missing = sorted(need - got)
        if missing:
            raise SheetError(f"The {name} tab does not have the {name} columns (missing: {', '.join(missing)}). "
                             "If this came from a published link, that tab is probably not published, "
                             "and Google sent a different tab instead.")


def current_counts():
    """Row counts in the data file the site is using now, or None if there is none."""
    if not OUT.exists():
        return None
    body = OUT.read_text(encoding="utf-8")
    start = body.find("{")
    try:
        cur = json.loads(body[start:body.rstrip().rstrip(";").rfind("}") + 1])
    except ValueError:
        return None
    return {k: len(cur.get(k, [])) for k in ("cards", "benefits", "earning")}


def main(tables, require=(), allow_shrink=False):
    missing = [t for t in require if t not in tables]
    if missing:
        raise SheetError(f"Missing tab(s): {', '.join(missing)}. Found: {sorted(tables)}")
    check_headers(tables)
    if "Cards" not in tables:
        raise SheetError(f"No 'Cards' tab. Tabs found: {sorted(tables)}")
    cards = build_cards(tables["Cards"])
    ids = {c["id"] for c in cards}
    benefits, dropped = ([], 0)
    if "Benefits" in tables:
        benefits, dropped = build_benefits(tables["Benefits"], ids)
    earning, e_dropped = ([], 0)
    if "Earning" in tables:
        earning, e_dropped = build_earning(tables["Earning"], ids)
    quarterly = build_quarterly(tables["Quarterly"], ids) if "Quarterly" in tables else []
    catalog = {"v": 1, "built": date.today().isoformat(), "cards": cards, "benefits": benefits,
               "earning": earning, "quarterly": quarterly}
    no_em_dashes(catalog)
    before = current_counts()
    if before and not allow_shrink:
        for k in ("cards", "benefits", "earning"):
            if before[k] and len(catalog[k]) < 0.8 * before[k]:
                raise SheetError(f"{k} would drop from {before[k]} to {len(catalog[k])} rows. "
                                 "That looks like a partial read, so nothing was written. "
                                 "If it is on purpose, run again with --allow-shrink.")
    body = json.dumps(catalog, ensure_ascii=False, indent=1)
    OUT.write_text(
        "/* GENERATED by tools/build_wallet_catalog.py from the My Wallet card catalog sheet.\n"
        "   Do not hand-edit. Change the sheet, download it, run the script. */\n"
        f"window.HRT_WALLET_CATALOG={body};\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUT.relative_to(OUT.parent.parent.parent)}: "
          f"{len(cards)} cards, {len(benefits)} benefits, {len(earning)} earning rates, "
          f"{len(quarterly)} quarters; left out as unverified: {dropped} benefits, {e_dropped} rates.")


if __name__ == "__main__":
    a = sys.argv[1:]
    require, allow_shrink = (), False
    if "--allow-shrink" in a:
        a.remove("--allow-shrink"); allow_shrink = True
    if "--require" in a:
        i = a.index("--require"); require = tuple(x for x in a[i + 1].split(",") if x); del a[i:i + 2]
    try:
        if len(a) == 2 and a[0] == "--csv-dir":
            t = tables_from_csv_dir(a[1])
        elif len(a) >= 2 and a[0] == "--urls":
            t = tables_from_urls(a[1:])
        elif len(a) == 1 and not a[0].startswith("--"):
            t = tables_from_xlsx(a[0])
        else:
            sys.exit(__doc__)
        main(t, require, allow_shrink)
    except SheetError as e:
        sys.exit(f"Not written. {e}")
